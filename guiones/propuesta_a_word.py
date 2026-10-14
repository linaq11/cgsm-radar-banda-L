# -*- coding: utf-8 -*-
"""Convierte Propuesta_trabajo_radar.md en .docx: carta, margenes 2,5 cm, Calibri 11,
justificado, titulos en negrita sin numerar."""
import io
import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(AQUI, "Propuesta_trabajo_radar.md")
SALIDA = os.path.join(AQUI, "Propuesta_trabajo_radar.docx")

d = docx.Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, lado, Cm(2.5))
est = d.styles["Normal"]
est.font.name = "Calibri"
est.font.size = Pt(11)
est.paragraph_format.space_after = Pt(6)
est.paragraph_format.line_spacing = 1.1


def corridas(p, texto, tam=11):
    for trozo in re.split(r"(\*\*.+?\*\*)", texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            r = p.add_run(trozo[2:-2]); r.bold = True
        else:
            r = p.add_run(trozo)
        r.font.size = Pt(tam)


for linea in io.open(MD, encoding="utf-8").read().split("\n"):
    t = linea.strip()
    if not t:
        continue
    if t.startswith("# "):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        corridas(p, "**" + t[2:] + "**", tam=14); p.paragraph_format.space_after = Pt(4)
    elif t.startswith("## "):
        p = d.add_paragraph(); corridas(p, "**" + t[3:] + "**", tam=11.5)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
    elif t.startswith("**Curso"):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; corridas(p, t, tam=10)
        p.paragraph_format.space_after = Pt(10)
    else:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; corridas(p, t)

d.save(SALIDA)
print("escrito:", SALIDA)
