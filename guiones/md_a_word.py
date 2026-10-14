# -*- coding: utf-8 -*-
"""Convierte un markdown del curso de radar a .docx: carta, margenes 2,5 cm, Calibri 11,
cuerpo justificado, encabezados en negrita sin numerar, tablas con encabezado sombreado.
El bloque de cabecera (lineas antes del primer '## ') va centrado, la ultima en negrita.

USO
    python md_a_word.py Avance_proyecto_radar.md
"""
import io
import os
import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ACENTO = "2E5090"
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def corridas(p, texto, tam=11, negrita=False, cursiva=False, color=None):
    for trozo in re.split(r"(\*\*.+?\*\*)", texto):
        if not trozo:
            continue
        neg, t = negrita, trozo
        if trozo.startswith("**"):
            neg, t = True, trozo[2:-2]
        r = p.add_run(t)
        r.font.size = Pt(tam)
        r.bold = neg or None
        r.italic = cursiva or None
        if color is not None:
            r.font.color.rgb = color


def parrafo(d, texto, tam=11, justificar=True, despues=6, negrita=False, cursiva=False, centrar=False):
    p = d.add_paragraph()
    corridas(p, texto, tam, negrita, cursiva)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrar else (WD_ALIGN_PARAGRAPH.JUSTIFY if justificar else WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_after = Pt(despues)
    return p


def sombrear(celda, hex_color):
    pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def ancho_figura(ruta):
    """Ancho en pulgadas: 6,3 como máximo, sin pasar de 5,5 de alto; las figuras casi cuadradas
    o verticales (gráficas de dispersión, matrices) se limitan a 5,0 de ancho."""
    from PIL import Image
    w, h = Image.open(ruta).size
    aspecto = h / w
    ancho = 6.3 if aspecto < 0.8 else 5.0
    return min(ancho, 5.5 / aspecto)


def margenes_celda(t, cm=0.08):
    pr = t._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for lado in ("left", "right"):
        e = OxmlElement("w:" + lado); e.set(qn("w:w"), str(int(cm * 567))); e.set(qn("w:type"), "dxa"); m.append(e)
    pr.append(m)


def tabla(d, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    celdas = [c for c in celdas if not all(re.fullmatch(r":?-+:?", x) for x in c)]
    enc, cuerpo = celdas[0], celdas[1:]
    tam = 9 if len(enc) <= 5 else (8 if len(enc) <= 7 else 7.5)
    t = d.add_table(rows=1, cols=len(enc))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    margenes_celda(t)
    for k, c in enumerate(enc):
        cel = t.rows[0].cells[k]; cel.text = ""
        corridas(cel.paragraphs[0], c, tam, negrita=True, color=BLANCO)
        cel.paragraphs[0].paragraph_format.space_after = Pt(0); sombrear(cel, ACENTO)
    tr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); tr.append(h)
    for fila in cuerpo:
        nueva = t.add_row().cells
        for k, c in enumerate(fila[:len(enc)]):
            nueva[k].text = ""; corridas(nueva[k].paragraphs[0], c, tam)
            nueva[k].paragraphs[0].paragraph_format.space_after = Pt(0)
            nueva[k].paragraphs[0].paragraph_format.line_spacing = 1.0
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def main(md):
    salida = os.path.splitext(md)[0] + ".docx"
    lineas = io.open(md, encoding="utf-8").read().split("\n")
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2.5))
    est = d.styles["Normal"]; est.font.name = "Calibri"; est.font.size = Pt(11); est.paragraph_format.line_spacing = 1.15
    primer = next(i for i, l in enumerate(lineas) if l.startswith("## "))
    cab = [l for l in lineas[:primer] if l.strip()]
    for l in cab[:-1]:
        parrafo(d, l, justificar=False, despues=0, centrar=True)
    parrafo(d, cab[-1], tam=13, negrita=True, despues=14, centrar=True)
    i, buf = primer, []
    while i < len(lineas):
        t = lineas[i].strip()
        if t.startswith("|"):
            buf.append(t); i += 1; continue
        if buf:
            tabla(d, buf); buf = []
        if t.startswith("#### "):
            p = parrafo(d, t[5:], tam=11, negrita=True, cursiva=True, justificar=False, despues=3)
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
        elif t.startswith("### "):
            p = parrafo(d, t[4:], tam=11, negrita=True, justificar=False, despues=3)
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
        elif t.startswith("## "):
            p = parrafo(d, t[3:], tam=12, negrita=True, justificar=False, despues=4)
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
        elif t.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", t)
            pie, ruta = m.group(1), os.path.join(os.path.dirname(os.path.abspath(md)), m.group(2))
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(ruta, width=Inches(ancho_figura(ruta)))
            p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
            parrafo(d, pie, tam=9.5, cursiva=True, despues=10)
        elif t.startswith("Tabla "):
            p = parrafo(d, t, tam=9.5, cursiva=True, justificar=False, despues=3, centrar=True); p.paragraph_format.keep_with_next = True
        elif t:
            parrafo(d, t)
        i += 1
    if buf:
        tabla(d, buf)
    d.save(salida)
    print("escrito:", salida)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Avance_proyecto_radar.md")
